"""
Tests for export REST API endpoints.

Covers:
- GET /export/formats
- GET /export/graph?format=json|graphml|gexf
- GET /export/matrices
- GET /export/matrix/{id}?format=xlsx|xmile
"""

import io
import json
import uuid
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from api.rest.dependencies import get_sfm_service
from models import Node
from models.delivery_matrix import Delivery
from graph.sfm_graph import Relationship


@pytest.fixture
def service(integration_app):
    return integration_app.dependency_overrides[get_sfm_service]()


@pytest.fixture
def populated(service):
    a = service.create_node(Node(label="Legislature", description="Funds"))
    b = service.create_node(Node(label="Districts", description="Spends"))
    service.create_relationship(Relationship(source_id=a.id, target_id=b.id, kind="funds", weight=0.8))
    matrix = service.create_delivery_matrix(label="Education Finance", description="Test matrix")
    matrix.add_component(a.id)
    matrix.add_component(b.id)
    service.add_delivery_to_matrix(
        matrix, a.id, b.id,
        Delivery(delivery_type="money", delivery_content="Appropriation", quantity=100.0, units="USD",
                 temporal_rate="annual"),
        cell_description="Legislature appropriates funds to districts",
    )
    return {"a": a, "b": b, "matrix": matrix}


class TestExportFormats:
    def test_lists_graph_and_matrix_formats(self, integration_client):
        resp = integration_client.get("/api/v1/export/formats")
        assert resp.status_code == 200
        formats = {f["format_name"]: f for f in resp.json()["formats"]}
        assert set(formats) == {"json", "graphml", "gexf", "xlsx", "xmile"}
        assert formats["json"]["scope"] == "graph"
        assert formats["xlsx"]["scope"] == "matrix"


class TestGraphExport:
    def test_json_export_downloads_snapshot(self, integration_client, populated):
        resp = integration_client.get("/api/v1/export/graph?format=json")
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/json")
        assert "sfm_graph.json" in resp.headers["content-disposition"]
        data = json.loads(resp.content)
        assert data["metadata"]["relationship_count"] == 1
        labels = {n["label"] for n in data["nodes"]}
        assert {"Legislature", "Districts"} <= labels

    def test_json_is_default_format(self, integration_client, populated):
        resp = integration_client.get("/api/v1/export/graph")
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/json")

    @pytest.mark.parametrize("fmt,root_tag", [("graphml", "graphml"), ("gexf", "gexf")])
    def test_xml_graph_formats_are_valid_xml(self, integration_client, populated, fmt, root_tag):
        resp = integration_client.get(f"/api/v1/export/graph?format={fmt}")
        assert resp.status_code == 200
        root = ET.fromstring(resp.content)
        assert root.tag.endswith(root_tag)
        assert f"sfm_graph.{fmt}" in resp.headers["content-disposition"]

    @pytest.mark.parametrize("fmt", ["graphml", "gexf"])
    def test_empty_graph_xml_export_is_400(self, integration_client, fmt):
        resp = integration_client.get(f"/api/v1/export/graph?format={fmt}")
        assert resp.status_code == 400
        assert "empty" in resp.json()["detail"].lower()

    def test_empty_graph_json_export_succeeds(self, integration_client):
        resp = integration_client.get("/api/v1/export/graph?format=json")
        assert resp.status_code == 200
        assert json.loads(resp.content)["metadata"]["node_count"] == 0

    def test_unknown_format_is_422(self, integration_client):
        resp = integration_client.get("/api/v1/export/graph?format=pdf")
        assert resp.status_code == 422


class TestMatrixExport:
    def test_list_matrices(self, integration_client, populated):
        resp = integration_client.get("/api/v1/export/matrices")
        assert resp.status_code == 200
        matrices = resp.json()["matrices"]
        assert len(matrices) == 1
        assert matrices[0]["id"] == str(populated["matrix"].id)
        assert matrices[0]["label"] == "Education Finance"
        assert matrices[0]["component_count"] == 2
        assert matrices[0]["cell_count"] == 1

    def test_xlsx_export_has_three_sheets(self, integration_client, populated):
        mid = populated["matrix"].id
        resp = integration_client.get(f"/api/v1/export/matrix/{mid}?format=xlsx")
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "Education_Finance.xlsx" in resp.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(resp.content))
        assert wb.sheetnames == ["Matrix View", "Cell Descriptions", "Delivery Details"]

    def test_xlsx_export_can_omit_optional_sheets(self, integration_client, populated):
        mid = populated["matrix"].id
        resp = integration_client.get(
            f"/api/v1/export/matrix/{mid}?format=xlsx"
            "&include_cell_descriptions=false&include_delivery_details=false"
        )
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        assert wb.sheetnames == ["Matrix View"]

    def test_xmile_export_is_valid_model(self, integration_client, populated):
        mid = populated["matrix"].id
        resp = integration_client.get(f"/api/v1/export/matrix/{mid}?format=xmile")
        assert resp.status_code == 200
        root = ET.fromstring(resp.content)
        assert root.tag.endswith("xmile")
        assert "Education_Finance.xmile" in resp.headers["content-disposition"]

    def test_xlsx_is_default_matrix_format(self, integration_client, populated):
        mid = populated["matrix"].id
        resp = integration_client.get(f"/api/v1/export/matrix/{mid}")
        assert resp.status_code == 200
        assert "xlsx" in resp.headers["content-disposition"]

    def test_unknown_matrix_is_404(self, integration_client, populated):
        resp = integration_client.get(f"/api/v1/export/matrix/{uuid.uuid4()}?format=xlsx")
        assert resp.status_code == 404

    def test_non_matrix_node_is_404(self, integration_client, populated):
        node_id = populated["a"].id
        resp = integration_client.get(f"/api/v1/export/matrix/{node_id}?format=xlsx")
        assert resp.status_code == 404

    def test_malformed_id_is_422(self, integration_client):
        resp = integration_client.get("/api/v1/export/matrix/not-a-uuid")
        assert resp.status_code == 422


class TestImportRouterUsesDependencyOverride:
    """The import router must go through Depends so tests can inject a service."""

    def test_csv_import_lands_in_overridden_service(self, integration_client, service):
        csv_bytes = b"name,description\nAlpha,First\nBeta,Second\n"
        resp = integration_client.post(
            "/api/v1/import/csv",
            files={"file": ("t.csv", io.BytesIO(csv_bytes), "text/csv")},
        )
        assert resp.status_code == 200
        assert resp.json()["nodes_created"] == 2
        assert {n.label for n in service.list_nodes()} == {"Alpha", "Beta"}
